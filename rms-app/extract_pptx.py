import sys
from pptx import Presentation

def extract_text_from_pptx(file_path):
    try:
        prs = Presentation(file_path)
    except Exception as e:
        print(f"Error opening file: {e}")
        return

    print(f"Total Slides: {len(prs.slides)}")
    print("-" * 20)

    for i, slide in enumerate(prs.slides):
        title = ""
        if slide.shapes.title:
            title = slide.shapes.title.text
        
        print(f"Slide {i+1}: {title if title else '[No Title]'}")
        
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip() and shape != slide.shapes.title:
                print(f"  - {shape.text.strip()}")
        print("-" * 20)

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python extract_pptx.py <path_to_pptx>")
    else:
        extract_text_from_pptx(sys.argv[1])
